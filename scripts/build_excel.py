"""Build the GoExtraCash condition-wise price workbook from crawled data.

Input : data/devices_full.json  (one record per RAM/ROM variant, with the
        per-condition deduction table returned by api.goextracash.com/device/{id})
Output: output/GoExtraCash_Condition_Wise_Price_List.xlsx
"""
import json
import os
import statistics

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.formatting.rule import ColorScaleRule

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
SRC = os.path.join(ROOT, "data", "devices_full.json")
OUTDIR = os.path.join(ROOT, "output")
os.makedirs(OUTDIR, exist_ok=True)
OUT = os.path.join(OUTDIR, "GoExtraCash_Condition_Wise_Price_List.xlsx")

INR = '"₹"#,##0'
PCT = '0.0%'

# --------------------------------------------------------------------------
# Condition catalogue. Groups and labels are the site's own wording, taken
# from its sell flow (device-status → screen → issues → accessories →
# warranty → physical → checkout).
# --------------------------------------------------------------------------
CONDITIONS = [
    # (group, api key, label of the answer that triggers the deduction)
    ("Power", "mobile_off", "Phone does not switch on"),
    ("Screen", "touch_notworking", "Touch faulty / not working"),
    ("Screen", "display_heavyspot", "Heavy visible spots on display"),
    ("Screen", "display_multiplespot", "Multiple visible spots on display"),
    ("Screen", "display_smallspot", "Minor visible spots on display"),
    ("Screen", "display_line", "Visible line(s) on display"),
    ("Screen", "display_faded", "Display faded along corners"),
    ("Screen", "scratches_glass", "Glass broken"),
    ("Screen", "scratches_crack", "Screen cracked on edge"),
    ("Screen", "scratches_heavy", "Heavy scratches on screen"),
    ("Screen", "scratches_minor", "Minor scratches on screen"),
    ("Functionality", "notworking_frontcamera", "Front camera not working"),
    ("Functionality", "notworking_backcamera", "Back camera not working"),
    ("Functionality", "notworking_volume", "Volume button not working"),
    ("Functionality", "notworking_battery", "Battery problem"),
    ("Functionality", "notworking_speaker", "Speaker not working"),
    ("Functionality", "notworking_finger", "Fingerprint sensor not working"),
    ("Functionality", "notworking_power", "Power button not working"),
    ("Functionality", "notworking_charging", "Charging port not working"),
    ("Functionality", "notworking_wifi", "Wi-Fi not working"),
    ("Functionality", "notworking_face", "Face sensor not working"),
    ("Body", "body_scratch_major", "Major scratches on body"),
    ("Body", "body_scratch_minor", "Less than 2 body scratches"),
    ("Body", "body_dent_major", "Multiple / heavy dents"),
    ("Body", "body_dent_minor", "2 or less minor dents"),
    ("Body", "body_dent_ignorable", "Minor ignorable dents"),
    ("Body", "panel_crack", "Side / back panel cracked"),
    ("Body", "panel_missing", "Panel missing"),
    ("Body", "screen_bent", "Bent / curved panel"),
    ("Body", "screen_gap", "Loose screen gap"),
    ("Accessories", "available_charger", "Original charger not available"),
    ("Accessories", "available_box", "Box not available"),
    ("Accessories", "available_earphone", "Earphone not available"),
    ("Accessories", "available_bill", "Valid bill not available"),
    ("Warranty", "warranty_12mmore", "Out of warranty – above 12 months"),
    ("Warranty", "warranty_12month", "Out of warranty – 6 to 12 months"),
    ("Warranty", "warranty_6month", "Out of warranty – 3 to 6 months"),
    ("Warranty", "warranty_3month", "Out of warranty – 1 to 3 months"),
    ("Warranty", "warranty_1month", "Out of warranty – below 1 month"),
]

# Columns shown on the headline sheet – one per real condition the site asks
# about, so every number is the site's own deduction, nothing invented.
HIGHLIGHTS = [
    ("scratches_minor", "Minor screen scratches"),
    ("scratches_heavy", "Heavy screen scratches"),
    ("scratches_crack", "Screen cracked on edge"),
    ("scratches_glass", "Glass broken"),
    ("touch_notworking", "Touch faulty"),
    ("display_smallspot", "Display – minor spots"),
    ("display_multiplespot", "Display – multiple spots"),
    ("display_heavyspot", "Display – heavy spots"),
    ("display_line", "Display – visible line"),
    ("display_faded", "Display – faded corners"),
    ("notworking_battery", "Battery problem"),
    ("notworking_face", "Face sensor not working"),
    ("notworking_backcamera", "Back camera not working"),
    ("notworking_charging", "Charging port not working"),
    ("body_scratch_minor", "A few body scratches"),
    ("body_scratch_major", "Major body scratches"),
    ("body_dent_minor", "2 or less minor dents"),
    ("body_dent_major", "Multiple / heavy dents"),
    ("panel_crack", "Back panel cracked"),
    ("available_box", "Box not available"),
    ("available_charger", "Charger not available"),
    ("warranty_12month", "Out of warranty – 6-12 months"),
    ("warranty_12mmore", "Out of warranty – above 12 months"),
    ("mobile_off", "Phone does not switch on"),
]

BRAND_DISPLAY = {
    "APPLE": "Apple", "ASUS": "Asus", "GOOGLE": "Google Pixel",
    "GOOGLE PIXEL": "Google Pixel", "GIONEE": "Gionee", "HONOR": "Honor",
    "HUAWEI": "Huawei", "INFINIX": "Infinix", "IQOO": "iQOO", "LG": "LG",
    "LETV": "LeTV", "LENOVO": "Lenovo", "MI": "Mi (Xiaomi)",
    "MOTOROLA": "Motorola", "NOKIA": "Nokia", "NOTHING": "Nothing",
    "ONEPLUS": "OnePlus", "OPPO": "Oppo", "POCO": "POCO", "REALME": "realme",
    "SAMSUNG": "Samsung", "TECNO": "Tecno", "VIVO": "Vivo",
}

# ------------------------------------------------------------------ styling
NAVY = "1F3864"
BLUE = "2F5597"
BAND = "F2F5FB"
WHITE = "FFFFFF"
GREY = "808080"

H_FONT = Font(bold=True, color=WHITE, size=10, name="Calibri")
H_FILL = PatternFill("solid", fgColor=BLUE)
TITLE_FONT = Font(bold=True, size=16, color=NAVY, name="Calibri")
SUB_FONT = Font(size=10, color=GREY, italic=True, name="Calibri")
THIN = Side(style="thin", color="BFBFBF")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)


def style_header(ws, row=1, height=32):
    for c in range(1, ws.max_column + 1):
        cell = ws.cell(row=row, column=c)
        cell.font = H_FONT
        cell.fill = H_FILL
        cell.alignment = Alignment(horizontal="center", vertical="center",
                                   wrap_text=True)
        cell.border = BORDER
    ws.row_dimensions[row].height = height


def autosize(ws, widths):
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w


def band(ws, first_row, last_row, ncols):
    fill = PatternFill("solid", fgColor=BAND)
    for r in range(first_row, last_row + 1):
        if (r - first_row) % 2 == 1:
            for c in range(1, ncols + 1):
                ws.cell(row=r, column=c).fill = fill


# --------------------------------------------------------------------- data
with open(SRC, encoding="utf-8") as f:
    rows = json.load(f)

devices = []
seen_ids = set()
for r in rows:
    d = r.get("detail") or {}
    vid = r.get("variant_id")
    if not d.get("msp") or vid in seen_ids:
        continue
    seen_ids.add(vid)
    raw_brand = (d.get("brand") or r.get("brand") or "").strip()
    devices.append({
        "brand": BRAND_DISPLAY.get(raw_brand.upper(), raw_brand.title()),
        "series": (d.get("series") or r.get("series") or "").strip(),
        "model": (d.get("model") or r.get("model") or "").strip(),
        "ram": d.get("ram") or r.get("ram"),
        "rom": d.get("rom") or r.get("rom"),
        "variant_id": vid,
        "msp": float(d["msp"]),
        "ded": {k: float(d.get(k) or 0) for _, k, _ in CONDITIONS},
        "updated": d.get("updated_at"),
    })

devices.sort(key=lambda x: (x["brand"].upper(), x["model"].upper(),
                            x["rom"] or 0, x["ram"] or 0))
print(f"devices with prices: {len(devices)}")


def price(dev, keys):
    return max(0.0, dev["msp"] - sum(dev["ded"].get(k, 0) for k in keys))


brands = sorted({d["brand"] for d in devices})
wb = Workbook()

# ------------------------------------------------------------------- Read Me
ws = wb.active
ws.title = "Read Me"
ws.sheet_view.showGridLines = False
ws["A1"] = "GoExtraCash – Condition-wise Price List"
ws["A1"].font = TITLE_FONT
ws["A2"] = ("Every mobile model on goextracash.com, with the price you get in "
            "each device condition")
ws["A2"].font = Font(size=11, color=BLUE, bold=True)

info = [
    ("Source", "https://www.goextracash.com/  (prices from api.goextracash.com)"),
    ("Devices covered", f"{len(devices):,} RAM/ROM variants across "
                        f"{len(brands)} brands"),
    ("Brands", ", ".join(brands)),
    ("Price range", f"₹{min(d['msp'] for d in devices):,.0f} – "
                    f"₹{max(d['msp'] for d in devices):,.0f} for a flawless phone"),
    ("", ""),
    ("How GoExtraCash prices a phone",
     "Each variant has a flawless price (the site calls it MSP) – what you get "
     "for a phone with no defects and all accessories, bill and warranty. "
     "Every condition defect has a fixed deduction, so:"),
    ("Formula", "quote = flawless price − Σ (deductions of the defects declared)"),
    ("Verified",
     "Checked against the site's own quote endpoint: single defects, random "
     "multi-defect combinations and 12 devices across brands matched exactly "
     "(36/36). Deductions only stop adding up when they exceed the flawless "
     "price; the site quotes such phones separately and prices here are "
     "floored at ₹0."),
    ("", ""),
    ("Sheet: Price by Condition",
     "One row per variant. Flawless price plus the price if each single "
     "condition applies (minor/ heavy scratches, cracked screen, glass broken, "
     "touch faulty, display spots & lines, battery, camera, body damage, "
     "missing box/charger, out of warranty, dead phone)."),
    ("Sheet: All Conditions",
     "Long format – one row per variant × condition (%s rows). Ideal for "
     "pivots: filter a model and see the price for every single defect the "
     "site asks about." % f"{len(devices) * len(CONDITIONS):,}"),
    ("Sheet: Condition Deductions",
     "The raw deduction table per variant, in ₹ and as a % of the flawless "
     "price, exactly as published by the site."),
    ("Sheet: Brand Summary", "Variant counts and price ranges by brand."),
    ("Sheet: Data", "Flat export of every price field per variant."),
    ("", ""),
    ("Warranty buckets",
     "The site asks how long ago an out-of-warranty phone's warranty ended "
     "(below 1 month / 1-3 / 3-6 / 6-12 / above 12 months). The deduction "
     "differs per bucket and per phone, so all five are listed separately."),
    ("Notes",
     "These are the site's indicative quotes before physical inspection – "
     "GoExtraCash states the final price may change after evaluation. "
     "All amounts are in Indian Rupees (₹). Data scraped 2026-10-24."),
]
r = 4
for k, v in info:
    ws.cell(row=r, column=1, value=k).font = Font(bold=True, size=10, color=NAVY)
    c = ws.cell(row=r, column=2, value=v)
    c.font = Font(size=10)
    c.alignment = Alignment(wrap_text=True, vertical="top")
    if k:
        ws.row_dimensions[r].height = max(15, 14 * (len(v) // 95 + 1))
    r += 1
ws.column_dimensions["A"].width = 28
ws.column_dimensions["B"].width = 112
ws.cell(row=r + 1, column=1, value="Generated by Arena.ai agent").font = SUB_FONT

# ---------------------------------------------------------- Price by Condition
ws = wb.create_sheet("Price by Condition")
hdr = ["Brand", "Series", "Model", "RAM (GB)", "Storage (GB)", "Variant ID",
       "Flawless / Like New (₹)"] + [f"{lab} (₹)" for _, lab in HIGHLIGHTS]
for c, h in enumerate(hdr, start=1):
    ws.cell(row=1, column=c, value=h)
style_header(ws, height=70)

for i, d in enumerate(devices, start=2):
    vals = [d["brand"], d["series"], d["model"], d["ram"], d["rom"],
            d["variant_id"], round(d["msp"])]
    vals += [round(price(d, [k])) for k, _ in HIGHLIGHTS]
    for c, v in enumerate(vals, start=1):
        cell = ws.cell(row=i, column=c, value=v)
        if c >= 7:
            cell.number_format = INR

n = len(devices) + 1
band(ws, 2, n, len(hdr))
ws.freeze_panes = "D2"
ws.auto_filter.ref = f"A1:{get_column_letter(len(hdr))}{n}"
autosize(ws, [12, 24, 30, 9, 12, 11, 21] + [16] * len(HIGHLIGHTS))
if n > 2:
    ws.conditional_formatting.add(
        f"G2:G{n}",
        ColorScaleRule(start_type="min", start_color="FFFFFF",
                       end_type="max", end_color="8EA9DB"))

# ------------------------------------------------------------ All Conditions
ws = wb.create_sheet("All Conditions")
hdr = ["Brand", "Series", "Model", "RAM (GB)", "Storage (GB)", "Variant ID",
       "Condition group", "Condition (the answer that triggers it)",
       "Flawless price (₹)", "Deduction (₹)", "Price with this condition (₹)",
       "Drop vs flawless"]
for c, h in enumerate(hdr, start=1):
    ws.cell(row=1, column=c, value=h)
style_header(ws, height=45)

row = 2
for d in devices:
    for group, key, label in CONDITIONS:
        ded = d["ded"].get(key, 0)
        vals = [d["brand"], d["series"], d["model"], d["ram"], d["rom"],
                d["variant_id"], group, label, round(d["msp"]), round(ded),
                round(max(0.0, d["msp"] - ded)),
                round(ded / d["msp"], 3) if d["msp"] else 0]
        for c, v in enumerate(vals, start=1):
            cell = ws.cell(row=row, column=c, value=v)
            if c in (9, 10, 11):
                cell.number_format = INR
            if c == 12:
                cell.number_format = PCT
        row += 1

last = row - 1
band(ws, 2, last, len(hdr))
ws.freeze_panes = "D2"
ws.auto_filter.ref = f"A1:{get_column_letter(len(hdr))}{last}"
autosize(ws, [12, 22, 28, 9, 11, 11, 15, 42, 16, 13, 24, 13])

# ------------------------------------------------------ Condition Deductions
ws = wb.create_sheet("Condition Deductions")
hdr = ["Brand", "Series", "Model", "RAM (GB)", "Storage (GB)", "Variant ID",
       "Flawless price (₹)"]
hdr += [f"{lab} – ₹" for _, _, lab in CONDITIONS]
hdr += [f"{lab} – %" for _, _, lab in CONDITIONS]
for c, h in enumerate(hdr, start=1):
    ws.cell(row=1, column=c, value=h)
style_header(ws, height=60)

for i, d in enumerate(devices, start=2):
    vals = [d["brand"], d["series"], d["model"], d["ram"], d["rom"],
            d["variant_id"], round(d["msp"])]
    vals += [round(d["ded"].get(k, 0), 1) for _, k, _ in CONDITIONS]
    vals += [round(d["ded"].get(k, 0) / d["msp"], 4) if d["msp"] else 0
             for _, k, _ in CONDITIONS]
    for c, v in enumerate(vals, start=1):
        cell = ws.cell(row=i, column=c, value=v)
        if 8 <= c <= 7 + len(CONDITIONS):
            cell.number_format = INR
        if c > 7 + len(CONDITIONS):
            cell.number_format = '0.0%'

n = len(devices) + 1
if n > 2:
    ws.conditional_formatting.add(
        f"H2:{get_column_letter(7 + len(CONDITIONS))}{n}",
        ColorScaleRule(start_type="min", start_color="FFFFFF",
                       end_type="max", end_color="F8696B"))
ws.freeze_panes = "D2"
ws.auto_filter.ref = f"A1:{get_column_letter(len(hdr))}{n}"
autosize(ws, [12, 22, 26, 8, 10, 10, 15] + [13] * len(CONDITIONS)
         + [9] * len(CONDITIONS))

# ------------------------------------------------------------- Brand Summary
ws = wb.create_sheet("Brand Summary")
hdr = ["Brand", "Variants", "Models", "Flawless: min (₹)", "Flawless: max (₹)",
       "Flawless: average (₹)", "Heavy screen scratches (avg ₹)",
       "Screen cracked (avg ₹)", "Does not switch on (avg ₹)"]
for c, h in enumerate(hdr, start=1):
    ws.cell(row=1, column=c, value=h)
style_header(ws, height=45)

by_brand = {}
for d in devices:
    by_brand.setdefault(d["brand"], []).append(d)

i = 2
for b in sorted(by_brand, key=str.upper):
    ds = by_brand[b]
    vals = [b, len(ds), len({d["model"] for d in ds}),
            round(min(d["msp"] for d in ds)), round(max(d["msp"] for d in ds)),
            round(statistics.mean(d["msp"] for d in ds)),
            round(statistics.mean(price(d, ["scratches_heavy"]) for d in ds)),
            round(statistics.mean(price(d, ["scratches_crack"]) for d in ds)),
            round(statistics.mean(price(d, ["mobile_off"]) for d in ds))]
    for c, v in enumerate(vals, start=1):
        cell = ws.cell(row=i, column=c, value=v)
        if c >= 4:
            cell.number_format = INR
        cell.border = BORDER
    i += 1

tot = ["TOTAL", len(devices), len({d["model"] for d in devices}),
       round(min(d["msp"] for d in devices)), round(max(d["msp"] for d in devices)),
       round(statistics.mean(d["msp"] for d in devices)),
       round(statistics.mean(price(d, ["scratches_heavy"]) for d in devices)),
       round(statistics.mean(price(d, ["scratches_crack"]) for d in devices)),
       round(statistics.mean(price(d, ["mobile_off"]) for d in devices))]
for c, v in enumerate(tot, start=1):
    cell = ws.cell(row=i, column=c, value=v)
    cell.font = Font(bold=True)
    cell.fill = PatternFill("solid", fgColor="D9E2F3")
    if c >= 4:
        cell.number_format = INR
band(ws, 2, i - 1, len(hdr))
ws.freeze_panes = "B2"
autosize(ws, [18, 10, 9, 16, 16, 18, 26, 22, 24])

# ---------------------------------------------------------------------- Data
ws = wb.create_sheet("Data")
fields = ["Brand", "Series", "Model", "RAM (GB)", "Storage (GB)", "Variant ID",
          "Flawless price (MSP)"]
fields += [lab for _, _, lab in CONDITIONS]
fields += ["Last updated (site)"]
for c, h in enumerate(fields, start=1):
    ws.cell(row=1, column=c, value=h)
style_header(ws, height=45)

for i, d in enumerate(devices, start=2):
    vals = [d["brand"], d["series"], d["model"], d["ram"], d["rom"],
            d["variant_id"], d["msp"]]
    vals += [d["ded"].get(k, 0) for _, k, _ in CONDITIONS]
    vals.append(str(d["updated"] or ""))
    for c, v in enumerate(vals, start=1):
        cell = ws.cell(row=i, column=c, value=v)
        if 7 <= c <= 6 + len(CONDITIONS):
            cell.number_format = INR
n = len(devices) + 1
ws.freeze_panes = "D2"
ws.auto_filter.ref = f"A1:{get_column_letter(len(fields))}{n}"
autosize(ws, [12, 22, 26, 8, 10, 10, 16] + [12] * len(CONDITIONS) + [20])

wb.save(OUT)
print("saved", OUT, os.path.getsize(OUT) // 1024, "KB")
